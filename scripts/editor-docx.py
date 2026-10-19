#!/usr/bin/env python3
"""Friday editor copy of next week's Week Ahead, as a Word document (Morgan 2026-10-19).

Reads the JSON that scripts/weekly-email.js writes with EDITOR_JSON=<path> and
builds a plain, easy-to-edit .docx: the intro, every public meeting (day by
day, with the summary the card shows) and every event with its suggested
language. John Metzger edits it; Morgan updates the newsletter on the Review
Desk from his edits. Sent by .github/workflows/editor-docx.yml.

Usage: python3 scripts/editor-docx.py editor.json Week-Ahead-2026-10-26.docx
"""
import json
import re
import sys
from datetime import date

from docx import Document
from docx.shared import Pt, RGBColor

GREY = RGBColor(0x5A, 0x6B, 0x64)


def day_label(iso):
    d = date.fromisoformat(iso)
    return d.strftime('%A, %B ') + str(d.day)


def add_marked(par, text):
    """Add text to a paragraph, turning **bold** marks into bold runs."""
    for seg in re.split(r'(\*\*[^*\n]+?\*\*)', text or ''):
        m = re.fullmatch(r'\*\*([^*\n]+?)\*\*', seg)
        if m:
            par.add_run(m.group(1)).bold = True
        elif seg:
            par.add_run(seg.replace('**', ''))


def strip_links(text):
    """[text](url) → text (url), so links survive as editable plain text."""
    return re.sub(r'\[([^\]]+)\]\((https?://[^)\s]+)\)', r'\1 (\2)', text or '')


def meta_line(doc, parts):
    p = doc.add_paragraph()
    r = p.add_run(' · '.join(x for x in parts if x))
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(2)
    return p


def build(data, out):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)

    doc.add_heading(data['subject'].replace(' — ', ': ', 1), level=0)
    note = doc.add_paragraph()
    note.add_run('Draft text for editing. ').bold = True
    note.add_run('This is everything in next week\'s Week Ahead email, in order: the intro, each public '
                 'meeting and each event, with the wording each one will carry. Edit anything directly '
                 '(Track Changes is welcome). Morgan will update the newsletter online from your edits '
                 'before it goes out Monday morning.')

    doc.add_heading('Intro', level=1)
    for para in [x.strip() for x in re.split(r'\n+|<p>', data.get('lede') or '') if x.strip()]:
        add_marked(doc.add_paragraph(), strip_links(para))

    meetings = data.get('meetings') or []
    doc.add_heading('Public Meetings This Week', level=1)
    if not meetings:
        doc.add_paragraph('No public meetings are listed for this week yet.')
    last = None
    for m in meetings:
        if m['date'] != last:
            doc.add_heading(day_label(m['date']), level=2)
            last = m['date']
        h = doc.add_paragraph()
        h.add_run(m['body']).bold = True
        h.paragraph_format.space_after = Pt(0)
        meta_line(doc, [m.get('time'), m.get('source')])
        add_marked(doc.add_paragraph(), strip_links(m.get('summary')))
        if m.get('whyThisMatters'):
            p = doc.add_paragraph()
            p.add_run('Why this matters: ').bold = True
            add_marked(p, strip_links(m['whyThisMatters']))
        if m.get('agenda'):
            meta_line(doc, ['Agenda: ' + m['agenda']])

    fests = data.get('festivals') or []
    if fests:
        doc.add_heading('Festivals', level=1)
        for f in fests:
            h = doc.add_paragraph()
            h.add_run(f['name']).bold = True
            h.paragraph_format.space_after = Pt(0)
            meta_line(doc, [f.get('dates'), f.get('url')])
            if f.get('promo'):
                add_marked(doc.add_paragraph(), strip_links(f['promo']))

    events = data.get('events') or []
    doc.add_heading('What to Attend (Events)', level=1)
    if not events:
        doc.add_paragraph('No events are listed for this week yet.')
    last = None
    for e in events:
        if e['date'] != last:
            doc.add_heading(day_label(e['date']), level=2)
            last = e['date']
        h = doc.add_paragraph()
        h.add_run(e['title']).bold = True
        h.paragraph_format.space_after = Pt(0)
        meta_line(doc, [e.get('time'), e.get('location')])
        add_marked(doc.add_paragraph(), strip_links(e.get('summary')))
        if e.get('url'):
            meta_line(doc, ['Details: ' + e['url']])

    doc.save(out)


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit('usage: editor-docx.py editor.json out.docx')
    with open(sys.argv[1], encoding='utf-8') as fh:
        build(json.load(fh), sys.argv[2])
    print('editor-docx: wrote ' + sys.argv[2])
