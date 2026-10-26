#!/usr/bin/env python3
"""Thursday-night Past Meetings draft, as a Word document (Morgan 2026-10-26).

Reads the JSON that scripts/past-meetings-editor.js writes and builds a plain,
easy-to-edit .docx: the intro, Major Decisions, Other Meetings, and meetings
whose recap isn't written yet. Emily edits it and sends it back; Morgan sends
it out on Friday. Sent by .github/workflows/past-meetings-docx.yml.

Usage: python3 scripts/past-meetings-docx.py past-editor.json Past-Meetings-2026-10-29.docx
"""
import importlib.util
import json
import os
import sys

from docx import Document
from docx.shared import Pt

# Same helpers and look as the Week Ahead editor copy.
_spec = importlib.util.spec_from_file_location('editor_docx', os.path.join(os.path.dirname(__file__), 'editor-docx.py'))
ed = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(ed)


def meeting(doc, m, with_text=True):
    h = doc.add_paragraph()
    h.add_run(m['body']).bold = True
    h.paragraph_format.space_after = Pt(0)
    ed.meta_line(doc, [ed.day_label(m['date']), m.get('time'), m.get('source')])
    if with_text:
        ed.add_marked(doc.add_paragraph(), ed.strip_links(m.get('text')))
        if m.get('link'):
            ed.meta_line(doc, ['Full summary: ' + m['link']])


def build(data, out):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)

    doc.add_heading(data['subject'].replace(' - ', ': ', 1), level=0)
    note = doc.add_paragraph()
    note.add_run('Draft text for editing. ').bold = True
    note.add_run('This is this week\'s Past Meetings email (' + data['label'] + '), in order: the intro, '
                 'the major decisions, the other meetings, and any meeting whose recap isn\'t written yet. '
                 'Edit anything directly (Track Changes is welcome) and send it back to Morgan, who will '
                 'send it out on Friday.')

    doc.add_heading('Intro', level=1)
    ed.add_marked(doc.add_paragraph(), data.get('lede'))

    if data.get('major'):
        doc.add_heading('Major Decisions', level=1)
        for m in data['major']:
            meeting(doc, m)

    if data.get('other'):
        doc.add_heading('Other Meetings' if data.get('major') else 'This Week\'s Meetings', level=1)
        for m in data['other']:
            meeting(doc, m)

    if not data.get('major') and not data.get('other'):
        doc.add_paragraph('No meeting recaps are available for this week yet.')

    if data.get('pending'):
        doc.add_heading('Recap Not Yet Available', level=1)
        p = doc.add_paragraph()
        r = p.add_run('These meetings were held this week but their video or minutes hadn\'t been summarized '
                      'by Thursday night. Add a summary or cut them.')
        r.italic = True
        for m in data['pending']:
            meeting(doc, m, with_text=False)

    doc.save(out)


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit('usage: past-meetings-docx.py past-editor.json out.docx')
    with open(sys.argv[1], encoding='utf-8') as fh:
        build(json.load(fh), sys.argv[2])
    print('past-meetings-docx: wrote ' + sys.argv[2])
